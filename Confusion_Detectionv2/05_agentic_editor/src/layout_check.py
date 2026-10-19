"""layout_check.py — finds slides whose layout is visibly broken.

Deterministic, no LLM. Three signals, each measured against the real
deck (6 slides known broken by eye: 5, 98, 99, 100, 278, 282):
  * characters drawn on top of other characters in the rendered PDF
    (flattened animations: several build steps printed at once);
  * characters lying inside an image's box in the PDF (a picture or
    diagram covering text);
  * in the .pptx, a paragraph-sized text box overlapping a picture or
    group, or another paragraph-sized text box.
No single signal finds all six without also flagging healthy slides
(a diagram drawn next to its own caption looks the same geometrically),
so this RANKS candidates for a person to confirm -- it does not decide.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Union

PROSE_MIN_CHARS = 40  # shorter text is a diagram label, not a paragraph
_CHAR_OVERLAP_SHARE = 0.3


@dataclass
class LayoutIssue:
    slide_id: int  # 0-indexed position, the same as Stage 1's slide_id for this deck
    score: float
    reasons: List[str] = field(default_factory=list)


def _rect(shape):
    return (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)


def _area(r) -> float:
    return max(r[2] - r[0], 0) * max(r[3] - r[1], 0)


def _intersection(a, b) -> float:
    return _area((max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])))


def _pdf_signals(pdf_path: Union[str, Path]) -> Dict[int, dict]:
    import pdfplumber

    # This deck's PDF has malformed colour operators; pdfminer logs one
    # warning per occurrence (~90) without affecting the text it reads.
    logging.getLogger("pdfminer").setLevel(logging.ERROR)
    signals = {}
    with pdfplumber.open(str(pdf_path)) as pdf:
        for index, page in enumerate(pdf.pages):
            chars = [c for c in page.chars if c["text"].strip()]
            grid: Dict[tuple, List[int]] = {}
            for i, c in enumerate(chars):
                grid.setdefault((int(c["x0"] // 20), int(c["top"] // 20)), []).append(i)
            overlapping = set()
            for (gx, gy), members in grid.items():
                near = [j for dx in (-1, 0, 1) for dy in (-1, 0, 1) for j in grid.get((gx + dx, gy + dy), [])]
                for i in members:
                    a = chars[i]
                    for j in near:
                        if j <= i:
                            continue
                        b = chars[j]
                        w = min(a["x1"], b["x1"]) - max(a["x0"], b["x0"])
                        h = min(a["bottom"], b["bottom"]) - max(a["top"], b["top"])
                        if w <= 0 or h <= 0:
                            continue
                        smaller = min((a["x1"] - a["x0"]) * (a["bottom"] - a["top"]), (b["x1"] - b["x0"]) * (b["bottom"] - b["top"])) or 1
                        double_drawn = a["text"] == b["text"] and abs(a["x0"] - b["x0"]) < 1 and abs(a["top"] - b["top"]) < 1  # bold/shadow
                        if w * h / smaller > _CHAR_OVERLAP_SHARE and not double_drawn:
                            overlapping.update((i, j))
            under_images = sum(
                1
                for c in chars
                for im in page.images
                if im["x0"] <= c["x0"] and c["x1"] <= im["x1"] and im["top"] <= c["top"] and c["bottom"] <= im["bottom"]
            )
            signals[index] = {"overlapping_chars": len(overlapping), "chars_under_images": under_images}
    return signals


def _pptx_signals(pptx_path: Union[str, Path]) -> Dict[int, dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx_path))
    slide_area = prs.slide_width * prs.slide_height
    signals = {}
    for index, slide in enumerate(prs.slides):
        shapes = list(slide.shapes)
        prose = [s for s in shapes if s.has_text_frame and len(s.text_frame.text.strip()) >= PROSE_MIN_CHARS]
        graphics = [s for s in shapes if s.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP) and _area(_rect(s)) < 0.95 * slide_area]
        covered = max((_intersection(_rect(t), _rect(g)) / (_area(_rect(t)) or 1) for t in prose for g in graphics), default=0.0)
        stacked = max(
            (
                _intersection(_rect(a), _rect(b)) / (min(_area(_rect(a)), _area(_rect(b))) or 1)
                for i, a in enumerate(prose)
                for b in prose[i + 1 :]
            ),
            default=0.0,
        )
        signals[index] = {"text_under_graphic": covered, "text_on_text": stacked}
    return signals


def find_layout_issues(pptx_path: Union[str, Path], pdf_path: Optional[Union[str, Path]] = None, max_results: int = 20) -> List[LayoutIssue]:
    """Candidates, most likely broken first. `pdf_path` (a rendering of
    the same deck) adds the two rendering-based signals; without it only
    the .pptx geometry is used."""
    geometry = _pptx_signals(pptx_path)
    rendering = _pdf_signals(pdf_path) if pdf_path is not None else {}
    issues = []
    for index, g in geometry.items():
        r = rendering.get(index, {"overlapping_chars": 0, "chars_under_images": 0})
        reasons, score = [], 0.0
        if r["overlapping_chars"] >= 10:
            reasons.append(f"{r['overlapping_chars']} characters are drawn on top of other text")
            score += 2.0 + min(r["overlapping_chars"] / 100, 2.0)
        if r["chars_under_images"] >= 20:
            reasons.append(f"{r['chars_under_images']} characters sit inside an image's area")
            score += 1.0 + min(r["chars_under_images"] / 300, 1.0)
        if g["text_under_graphic"] > 0.2:
            reasons.append(f"a picture or diagram overlaps {g['text_under_graphic']:.0%} of a paragraph")
            score += g["text_under_graphic"]
        if g["text_on_text"] > 0.25:
            reasons.append(f"two paragraphs overlap by {g['text_on_text']:.0%}")
            score += g["text_on_text"]
        if reasons:
            issues.append(LayoutIssue(slide_id=index, score=round(score, 3), reasons=reasons))
    issues.sort(key=lambda i: -i.score)
    return issues[:max_results]
