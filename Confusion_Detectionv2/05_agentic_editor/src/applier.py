"""applier.py — applies accepted SlideEdits to a copy of the deck.

The source .pptx is never written to. Three operations:
  * insert_slide_after: a new slide after the anchor;
  * append_bullets: bullets added to the end of the anchor's own body,
    formatted like its last bullet;
  * repair_slide: a clean rebuild of a broken slide right after it --
    its own pictures moved into a panel, its text as bullets -- with the
    original hidden from the slideshow (kept, never deleted).
New slides borrow their header (the course banner + title placeholder
and the rule under it) from the nearest existing slide that has one, so
they match the deck's template without that template being hardcoded.
"""
from __future__ import annotations

import copy
import math
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .checker import duplicate_anchor_ids
from .schema import APPEND_BULLETS, INSERT_SLIDE_AFTER, REPAIR_SLIDE, AppliedEdit, SlideEdit
from .slide_parts import deck_banner, split_slide

ADDED_SLIDE_NAME = "gap_editor_added"
BODY_SHAPE_NAME = "gap_editor_body"
BODY_FONT, BODY_FONT_PT = "Calibri", 22
MIN_FIT_FONT_PT = 12
# Body box position/size of a typical content slide in a 13.33in-wide
# deck, expressed as fractions so other slide sizes scale with it.
_BODY_LEFT, _BODY_TOP, _BODY_WIDTH, _BODY_HEIGHT = 0.036, 0.216, 0.92, 0.70
_RULE_MAX_HEIGHT = Inches(0.1)
# 128 of the deck's headers have a 12.8 in title box; new slides copy one
# of those, never one of the few narrow ones.
_MIN_HEADER_WIDTH = Inches(8)
# Repaired slides: shares of the body area tried for the pictures, when
# they sit to the right of the text or below it. Text-heavy slides
# (278, 282) need the smaller shares to stay readable.
_PICTURE_SHARES_RIGHT, _PICTURE_SHARES_BELOW = (0.38, 0.48), (0.45, 0.55)
_PANEL_GAP = 0.02  # of the slide width/height
_MAX_PICTURE_SCALE = 1.5  # a small picture is enlarged, but not blown up
_MIN_LABEL_SZ = 900  # 9pt, in hundredths of a point
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _header_shapes(slide):
    """(title placeholder, rule) of a slide that follows the deck's
    content template -- a banner line plus a title line, with a rule
    under them -- else None. A slide with a title placeholder but not
    this header (the real deck has some: a one-line title holding a
    whole sentence, no banner, no rule) is not a template to copy."""
    title = slide.shapes.title
    if title is None or not title.has_text_frame:
        return None
    if len([p for p in title.text_frame.paragraphs if p.text.strip()]) < 2:
        return None
    if title.width < _MIN_HEADER_WIDTH:
        return None  # a narrow title box (a few slides have 3-6 in) wraps a new title under the rule
    rule = next(
        (s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.FREEFORM and s.height <= _RULE_MAX_HEIGHT and s.top < title.top + title.height + Inches(1)),
        None,
    )
    return (title, rule) if rule is not None else None


def _nearest_template_slide(prs, position: int):
    slides = prs.slides
    for distance in range(len(slides)):
        for candidate in (position - distance, position + distance):
            if 0 <= candidate < len(slides) and _header_shapes(slides[candidate]) is not None:
                return slides[candidate]
    return None


def _set_title(title_shape, text: str) -> None:
    """The deck's title placeholder holds the course banner on its
    first line and the slide title on its last; only the last non-empty
    paragraph is rewritten, keeping its first run's formatting."""
    paragraphs = [p for p in title_shape.text_frame.paragraphs if p.text.strip()]
    target = paragraphs[-1]
    target.runs[0].text = plain_text(text)
    # Most of the deck's titles use the layout's default size; a size
    # copied from the template slide (some shrink long titles) made new
    # titles differ from slide to slide.
    r_pr = target.runs[0]._r.find(qn("a:rPr"))
    if r_pr is not None and "sz" in r_pr.attrib:
        del r_pr.attrib["sz"]
    for run in target.runs[1:]:
        run._r.getparent().remove(run._r)


def _add_bullet(paragraph, text: str, size_pt: float = BODY_FONT_PT) -> None:
    run = paragraph.add_run()
    run.text = plain_text(text)
    run.font.name, run.font.size = BODY_FONT, Pt(size_pt)
    paragraph.space_after = Pt(round(size_pt * 0.45))
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.35)))
    pPr.set("indent", str(-Inches(0.35)))
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu_font)
    pPr.append(bu_char)


# Characters LLMs like that Calibri cannot draw (PowerPoint shows a box
# or swaps the font): non-breaking hyphens and spaces become plain ones.
_PLAIN = str.maketrans({"\u2010": "-", "\u2011": "-", "\u202f": " ", "\u00a0": " ", "\u2007": " "})


def plain_text(text: str) -> str:
    return text.translate(_PLAIN)


def fit_font_pt(bullets: List[str], width_emu: int, height_emu: int, largest: float = BODY_FONT_PT) -> float:
    """The largest font size (down to MIN_FIT_FONT_PT) at which the
    bullets fit the box, by an average-character-width estimate
    (Calibri averages a little under half its point size per character)."""
    width_pt = (width_emu / 914400) * 72 - 0.35 * 72  # less the bullet indent
    height_pt = (height_emu / 914400) * 72
    size = largest
    while size > MIN_FIT_FONT_PT:
        per_line = max(int(width_pt / (size * 0.46)), 1)
        lines = sum(max(math.ceil(len(b) / per_line), 1) for b in bullets)
        needed = lines * size * 1.15 + (len(bullets) - 1) * size * 0.45
        if needed <= height_pt:
            return size
        size -= 1
    return MIN_FIT_FONT_PT


def _new_slide_with_header(prs, template_slide, title: str):
    """Appends a slide at the end of the deck carrying the template's
    header with `title`, and returns it."""
    layout = template_slide.slide_layout if template_slide is not None else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):  # the layout's empty placeholders
        shape._element.getparent().remove(shape._element)

    header = _header_shapes(template_slide) if template_slide is not None else None
    if header is not None:
        for source in header:
            slide.shapes._spTree.append(copy.deepcopy(source._element))
        _set_title(slide.shapes.title, title)
    else:
        box = slide.shapes.add_textbox(Inches(0.3), Inches(0.3), prs.slide_width - Inches(0.6), Inches(0.9))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = title
        run.font.size, run.font.bold = Pt(24), True
    slide._element.find(qn("p:cSld")).set("name", ADDED_SLIDE_NAME)
    return slide


def deck_bullet_style(prs, position: int = 0):
    """(paragraph properties, run properties, size in pt) of the deck's
    MOST COMMON bullet style, copied from a real paragraph that has it,
    so every new slide looks like the typical slide. In this deck: a
    round Arial bullet, 22pt, normal weight, default colour (134
    paragraphs); dark-orange squares appear in only some lectures. The
    nearest slide's style was tried first and gave new slides bold, blue
    or tiny text whenever the neighbour was unusual. `position` is kept
    for callers; the choice does not depend on it. None if the deck has
    no bullets."""
    from collections import Counter

    cache = getattr(prs, "_gap_editor_bullet_style", None)
    if cache is not None:
        return cache[0]
    examples, counts = {}, Counter()
    for slide in prs.slides:
        title = slide.shapes.title
        for shape in slide.shapes:
            if not shape.has_text_frame or (title is not None and shape.shape_id == title.shape_id) or shape.name == BODY_SHAPE_NAME:
                continue
            for paragraph in shape.text_frame.paragraphs:
                p_pr = paragraph._p.find(qn("a:pPr"))
                runs = paragraph._p.findall(qn("a:r"))
                if p_pr is None or p_pr.find(qn("a:buChar")) is None or not runs or len(paragraph.text.strip()) < 25:
                    continue
                r_pr = runs[0].find(qn("a:rPr"))
                if r_pr is None or not r_pr.get("sz") or r_pr.get("b") == "1" or r_pr.find(qn("a:solidFill")) is not None:
                    continue  # bold or coloured text is an emphasis, not the body style
                font = p_pr.find(qn("a:buFont"))
                colour = p_pr.find(qn("a:buClr"))
                key = (
                    p_pr.find(qn("a:buChar")).get("char"),
                    font.get("typeface") if font is not None else None,
                    colour.find(qn("a:srgbClr")).get("val") if colour is not None and colour.find(qn("a:srgbClr")) is not None else None,
                    r_pr.get("sz"),
                )
                counts[key] += 1
                examples.setdefault(key, (copy.deepcopy(p_pr), copy.deepcopy(r_pr)))
    style = None
    if counts:
        key = counts.most_common(1)[0][0]
        p_pr, r_pr = examples[key]
        style = (p_pr, r_pr, int(key[3]) / 100)
    prs._gap_editor_bullet_style = (style,)
    return style


def _add_styled_bullet(paragraph, text: str, style, size_pt: float) -> None:
    p_pr, r_pr, _ = style
    old = paragraph._p.find(qn("a:pPr"))
    if old is not None:
        paragraph._p.remove(old)
    paragraph._p.insert(0, copy.deepcopy(p_pr))
    run = paragraph.add_run()
    run.text = plain_text(text)
    new_rpr = copy.deepcopy(r_pr)
    new_rpr.set("sz", str(int(size_pt * 100)))
    old_rpr = run._r.find(qn("a:rPr"))
    if old_rpr is not None:
        run._r.remove(old_rpr)
    run._r.insert(0, new_rpr)


def _add_body(slide, bullets: List[str], rect: Tuple[int, int, int, int], size_pt: float, style=None):
    left, top, right, bottom = rect
    body = slide.shapes.add_textbox(left, top, right - left, bottom - top)
    body.name = BODY_SHAPE_NAME
    frame = body.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if style is not None:
            _add_styled_bullet(paragraph, bullet, style, size_pt)
        else:
            _add_bullet(paragraph, bullet, size_pt)
    return body


def _body_rect(prs) -> Tuple[int, int, int, int]:
    width, height = prs.slide_width, prs.slide_height
    left, top = int(width * _BODY_LEFT), int(height * _BODY_TOP)
    return (left, top, left + int(width * _BODY_WIDTH), top + int(height * _BODY_HEIGHT))


def _build_slide(prs, template_slide, edit: SlideEdit, style=None):
    """Appends the new slide at the end of the deck and returns it. With
    `style` (deck_bullet_style), its bullets copy the deck's own; the
    font shrinks from the deck's size only if the text would not fit."""
    slide = _new_slide_with_header(prs, template_slide, edit.title)
    rect = _body_rect(prs)
    size = style[2] if style else BODY_FONT_PT
    size = fit_font_pt(edit.bullets, rect[2] - rect[0], rect[3] - rect[1], largest=size)
    _add_body(slide, edit.bullets, rect, size, style)
    answered = "; ".join(edit.addresses_questions)
    slide.notes_slide.notes_text_frame.text = (
        f"Added by the gap editor for topic {edit.topic_id} ({edit.gap_type}). Student questions answered: {answered}"
    )
    return slide


# ---- append_bullets ------------------------------------------------------


def _body_shape(slide):
    """The slide's main text: its largest non-title text box with text."""
    title = slide.shapes.title
    candidates = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame
        and shape.text_frame.text.strip()
        and (title is None or shape.shape_id != title.shape_id)
    ]
    return max(candidates, key=lambda shape: shape.width * shape.height, default=None)


def _wrap_width(slide, body) -> Optional[int]:
    """Width the new text must keep to when a picture sits on top of the
    body text box, as on slide 137 (diagram over the box's right half):
    the slide's own bullets avoid it with manual line breaks, and added
    bullets must do the same. None when nothing covers the box."""
    left, right = body.left, body.left + body.width
    edges = [
        shape.left
        for shape in slide.shapes
        if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP)
        and left < shape.left < right
        and shape.top < body.top + body.height
        and shape.top + shape.height > body.top
    ]
    return min(edges) - left - Inches(0.2) if edges else None


def _break_lines(text: str, width_emu: int, size_pt: float) -> List[str]:
    """Lines of at most the characters that fit the width (half the font
    size per character: conservative for Calibri)."""
    import textwrap

    per_line = max(int((width_emu / 914400) * 72 / (size_pt * 0.5)), 12)
    return textwrap.wrap(text, per_line) or [text]


def _append_bullets(shape, bullets: List[str], wrap_width: Optional[int] = None) -> None:
    """Each new paragraph is a copy of the body's last non-empty one --
    same level, bullet, font and size -- with the new text in its first
    run. With `wrap_width`, the text is split into lines joined by line
    breaks so it stays clear of a picture over the box."""
    paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    template = paragraphs[-1]._p
    anchor = template
    for text in bullets:
        new_p = copy.deepcopy(template)
        runs = new_p.findall(qn("a:r"))
        for extra in runs[1:] + new_p.findall(qn("a:br")) + new_p.findall(qn("a:fld")):
            new_p.remove(extra)
        first = runs[0]
        rpr = first.find(qn("a:rPr"))
        size_pt = int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else 18.0
        text = plain_text(text)
        lines = _break_lines(text, wrap_width, size_pt) if wrap_width else [text]
        first.find(qn("a:t")).text = lines[0]
        previous = first
        for line in lines[1:]:
            br = new_p.makeelement(qn("a:br"), {})
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            run = copy.deepcopy(first)
            run.find(qn("a:t")).text = line
            previous.addnext(br)
            br.addnext(run)
            previous = run
        anchor.addnext(new_p)
        anchor = new_p


# ---- room for added bullets -----------------------------------------------

_BOTTOM_MARGIN = Inches(0.35)


def _paragraph_lines(text: str, width_emu: int, size_pt: float) -> int:
    per_line = max(int((width_emu / 914400) * 72 / (size_pt * 0.46)), 1)
    return sum(max(math.ceil(len(piece) / per_line), 1) for piece in text.split("\v"))


def _paragraph_size(paragraph) -> float:
    sizes = [r.font.size.pt for r in paragraph.runs if r.font.size]
    return sizes[0] if sizes else 18.0


def _text_width(body, paragraph) -> int:
    p_pr = paragraph._p.find(qn("a:pPr"))
    margins = sum(int(p_pr.get(k, 0)) for k in ("marL", "marR")) if p_pr is not None else 0
    body_pr = body.text_frame._txBody.find(qn("a:bodyPr"))
    insets = sum(int(body_pr.get(k, 91440)) for k in ("lIns", "rIns")) if body_pr is not None else 2 * 91440
    return max(body.width - margins - insets, Inches(1))


def append_room_problem(slide, slide_height: int, bullets: List[str]) -> Optional[str]:
    """Why the bullets would not fit below the slide's own text, or None.
    Estimates where the existing text ends (its real font sizes, margins
    and line breaks; ~0.46 x font size per character) and how much the
    new bullets add, formatted like the last bullet."""
    body = _body_shape(slide)
    if body is None:
        return "has no body text box to add to"
    paragraphs = [p for p in body.text_frame.paragraphs if p.text.strip()]
    body_pr = body.text_frame._txBody.find(qn("a:bodyPr"))
    y = body.top + (int(body_pr.get("tIns", 45720)) if body_pr is not None else 45720)
    for p in paragraphs:
        size = _paragraph_size(p)
        y += int(_paragraph_lines(p.text, _text_width(body, p), size) * size * 1.2 * 12700)
    last = paragraphs[-1]
    size = _paragraph_size(last)
    needed = sum(int(_paragraph_lines(b, _text_width(body, last), size) * size * 1.2 * 12700) for b in bullets)
    limit = slide_height - _BOTTOM_MARGIN
    if y + needed > limit:
        free = max(limit - y, 0) / 914400
        return f"has no room: its text already reaches {y / 914400:.1f} in of {slide_height / 914400:.1f} in, leaving {free:.1f} in; the new text needs {needed / 914400:.1f} in"
    return None


class RoomChecker:
    """Answers "do these bullets fit on slide N?" from the real .pptx,
    for the checker (which only sees the parsed text)."""

    def __init__(self, pptx_path):
        self._prs = Presentation(str(pptx_path))

    def __call__(self, edit: SlideEdit) -> Optional[str]:
        if edit.operation != APPEND_BULLETS or not 0 <= edit.after_slide_id < len(self._prs.slides):
            return None
        problem = append_room_problem(self._prs.slides[edit.after_slide_id], self._prs.slide_height, edit.bullets)
        return f"slide {edit.after_slide_id + 1} {problem}; make it a new slide instead" if problem else None


def _note(slide, line: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.text = f"{frame.text}\n{line}".strip()


# ---- repair_slide --------------------------------------------------------


def _set_xfrm(element, rect: Tuple[int, int, int, int]) -> None:
    tag = element.tag.split("}")[1]
    if tag == "grpSp":
        xfrm = element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    elif tag == "graphicFrame":
        xfrm = element.find(qn("p:xfrm"))
    else:
        xfrm = element.find(qn("p:spPr")).find(qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("x", str(rect[0]))
    xfrm.find(qn("a:off")).set("y", str(rect[1]))
    xfrm.find(qn("a:ext")).set("cx", str(max(rect[2] - rect[0], 0)))
    xfrm.find(qn("a:ext")).set("cy", str(max(rect[3] - rect[1], 0)))


def _relink(element, source_part, target_part) -> None:
    """Points every relationship the copied element uses (pictures,
    links) at the same target from the new slide's part."""
    for node in element.iter():
        for attribute, value in list(node.attrib.items()):
            if not attribute.startswith("{" + _R_NS + "}") or value not in source_part.rels:
                continue
            rel = source_part.rels[value]
            if rel.is_external:
                new_id = target_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_id = target_part.relate_to(rel.target_part, rel.reltype)
            node.set(attribute, new_id)


def _scale_fonts(element, scale: float) -> None:
    if scale >= 1:
        return
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for node in element.iter(qn(tag)):
            if node.get("sz"):
                node.set("sz", str(max(int(int(node.get("sz")) * scale), _MIN_LABEL_SZ)))


def _renumber(elements, slide) -> None:
    """Gives copied shapes ids unused on the new slide (duplicate ids make
    PowerPoint report the file as damaged), keeping connectors attached."""
    next_id = 1 + max((int(n.get("id")) for n in slide.shapes._spTree.iter(qn("p:cNvPr"))), default=1)
    mapping: Dict[str, str] = {}
    for element in elements:
        for node in element.iter(qn("p:cNvPr")):
            mapping.setdefault(node.get("id"), str(next_id))
            node.set("id", str(next_id))
            next_id += 1
    for element in elements:
        for tag in ("a:stCxn", "a:endCxn"):
            for node in element.iter(qn(tag)):
                if node.get("id") in mapping:
                    node.set("id", mapping[node.get("id")])


def _union(rects):
    return (min(r[0] for r in rects), min(r[1] for r in rects), max(r[2] for r in rects), max(r[3] for r in rects))


def _repair_layout(prs, bullets, graphics_rect):
    """(text rect, picture rect, font size, picture scale) for a repaired
    slide: pictures right of the text or below it, in the arrangement
    that keeps the text largest (16pt, else 14pt) and, among those, shows
    the pictures largest."""
    body = _body_rect(prs)
    bw, bh = body[2] - body[0], body[3] - body[1]
    gap_w, gap_h = int(prs.slide_width * _PANEL_GAP), int(prs.slide_height * _PANEL_GAP)
    uw, uh = graphics_rect[2] - graphics_rect[0], graphics_rect[3] - graphics_rect[1]
    options = []
    for share in _PICTURE_SHARES_RIGHT:
        split = body[0] + int(bw * (1 - share))
        options.append(((body[0], body[1], split - gap_w, body[3]), (split, body[1], body[2], body[3])))
    for share in _PICTURE_SHARES_BELOW:
        split = body[1] + int(bh * (1 - share))
        options.append(((body[0], body[1], body[2], split - gap_h), (body[0], split, body[2], body[3])))
    scored = []
    for text_rect, picture_rect in options:
        pw, ph = picture_rect[2] - picture_rect[0], picture_rect[3] - picture_rect[1]
        scale = min(pw / max(uw, 1), ph / max(uh, 1), _MAX_PICTURE_SCALE)
        size = fit_font_pt(bullets, text_rect[2] - text_rect[0], text_rect[3] - text_rect[1])
        scored.append(((size >= 16, size >= 14), scale, size, text_rect, picture_rect))
    _, scale, size, text_rect, picture_rect = max(scored, key=lambda o: (o[0], o[1], o[2]))
    return text_rect, picture_rect, size, scale


def _build_repaired_slide(prs, template_slide, source_slide, banner, edit: SlideEdit, style=None):
    """Appends the rebuild of `source_slide` at the end of the deck."""
    parts = split_slide(source_slide, prs.slide_width, prs.slide_height, banner)
    slide = _new_slide_with_header(prs, template_slide, edit.title)
    graphics = parts.graphics if edit.keep_pictures else []
    if not graphics:
        largest = style[2] if style else BODY_FONT_PT
        size = fit_font_pt(edit.bullets, _body_rect(prs)[2] - _body_rect(prs)[0], _body_rect(prs)[3] - _body_rect(prs)[1], largest=largest)
        _add_body(slide, edit.bullets, _body_rect(prs), size, style)
    else:
        union = _union([g.rect for g in graphics])
        text_rect, panel, size, scale = _repair_layout(prs, edit.bullets, union)
        _add_body(slide, edit.bullets, text_rect, min(size, style[2]) if style else size, style)
        # Centre the scaled pictures in their panel.
        used_w, used_h = (union[2] - union[0]) * scale, (union[3] - union[1]) * scale
        x0 = panel[0] + ((panel[2] - panel[0]) - used_w) / 2
        y0 = panel[1] + ((panel[3] - panel[1]) - used_h) / 2
        copies = []
        for g in graphics:
            element = copy.deepcopy(g.element)
            r = g.rect
            _set_xfrm(element, (
                int(x0 + (r[0] - union[0]) * scale),
                int(y0 + (r[1] - union[1]) * scale),
                int(x0 + (r[2] - union[0]) * scale),
                int(y0 + (r[3] - union[1]) * scale),
            ))
            _scale_fonts(element, scale)
            _relink(element, source_slide.part, slide.part)
            copies.append(element)
        _renumber(copies, slide)
        for element in copies:
            slide.shapes._spTree.append(element)
    return slide


def _move_last_slide_to(prs, position: int) -> None:
    slide_ids = prs.slides._sldIdLst
    element = slide_ids[-1]
    slide_ids.remove(element)
    slide_ids.insert(position, element)


def apply_edits(source_pptx: Union[str, Path], output_pptx: Union[str, Path], edits: List[SlideEdit], deck) -> List[AppliedEdit]:
    """Writes output_pptx = source_pptx with every edit applied. `deck`
    is Stage 1's parse of the same deck: its slide_ids are what
    after_slide_id refers to, and its row order is the slide order in
    the file. Inserted slides and added bullets also go into verbatim
    repeats of their anchor (the duplicated lecture); a repair applies to
    the one slide a person confirmed."""
    source_pptx, output_pptx = Path(source_pptx), Path(output_pptx)
    if source_pptx.resolve() == output_pptx.resolve():
        raise ValueError("output_pptx must differ from source_pptx: the source deck is never modified")

    prs = Presentation(str(source_pptx))
    if len(prs.slides) != len(deck.slides):
        raise ValueError(f"{source_pptx.name} has {len(prs.slides)} slides but the parsed deck has {len(deck.slides)}: slide ids would not line up")
    position_of = {s.slide_id: i for i, s in enumerate(deck.slides)}
    banner = deck_banner(prs)
    original = list(prs.slides)  # by original position, unaffected by insertions

    mirrored, notes = {}, {}
    insertions = []  # (anchor position, edit index): a new slide right after the anchor
    appends = {}  # edit index -> anchor positions changed in place
    for index, edit in enumerate(edits):
        mirrored[index] = [] if edit.operation == REPAIR_SLIDE else duplicate_anchor_ids(deck, edit.after_slide_id)
        anchors = [position_of[a] for a in [edit.after_slide_id] + mirrored[index]]
        if edit.operation == APPEND_BULLETS:
            if any(_body_shape(original[p]) is None for p in anchors):
                notes[index] = "the slide has no body text box to add to, so the bullets went on a new slide after it"
                insertions += [(p, index) for p in anchors]
            else:
                appends[index] = anchors
        else:
            insertions += [(p, index) for p in anchors]
    insertions.sort()

    # In place first, while positions are the original ones.
    for index, anchors in appends.items():
        edit = edits[index]
        for p in anchors:
            body = _body_shape(original[p])
            _append_bullets(body, edit.bullets, _wrap_width(original[p], body))
            _note(original[p], f"Gap editor: added the last {len(edit.bullets)} bullet(s) for topic {edit.topic_id} ({edit.gap_type}). "
                                f"Student questions answered: {'; '.join(edit.addresses_questions)}")
    for index, edit in enumerate(edits):
        if edit.operation == REPAIR_SLIDE:
            original[position_of[edit.after_slide_id]]._element.set("show", "0")

    # Then insertions, back to front, so positions not yet processed are
    # still the original ones.
    for anchor_position, index in reversed(insertions):
        edit = edits[index]
        template = _nearest_template_slide(prs, anchor_position)
        style = deck_bullet_style(prs, anchor_position)
        if edit.operation == REPAIR_SLIDE:
            new = _build_repaired_slide(prs, template, original[anchor_position], banner, edit, style)
            new.notes_slide.notes_text_frame.text = (
                f"Rebuilt by the gap editor from slide {anchor_position + 1} of the original deck, whose layout was broken: "
                "its own text restructured, its own pictures moved. The original is kept, hidden, right before this slide."
            )
        elif edit.operation == APPEND_BULLETS:  # fallback: no body text box on the anchor
            _build_slide(prs, template, SlideEdit(**{**edit.__dict__, "title": f"{edit.title} (continued)".strip()}), style)
        else:
            _build_slide(prs, template, edit, style)
        _move_last_slide_to(prs, anchor_position + 1)

    def final(position: int) -> int:
        """Where an original slide ends up: shifted by every insertion
        after an earlier slide."""
        return position + sum(1 for p, _ in insertions if p < position)

    positions = {index: [] for index in range(len(edits))}
    for rank, (anchor_position, index) in enumerate(insertions):
        positions[index].append(anchor_position + 1 + rank)
    for index, anchors in appends.items():
        positions[index] = [final(p) for p in anchors]

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))
    return [
        AppliedEdit(
            edit=edit,
            new_slide_positions=positions[index],
            mirrored_after_slide_ids=mirrored[index],
            hidden_slide_positions=[final(position_of[edit.after_slide_id])] if edit.operation == REPAIR_SLIDE else [],
            note=notes.get(index),
        )
        for index, edit in enumerate(edits)
    ]


def added_slide_positions(pptx_path: Union[str, Path]) -> List[int]:
    """0-indexed positions of the slides this module added to a deck."""
    prs = Presentation(str(pptx_path))
    return [i for i, s in enumerate(prs.slides) if s._element.find(qn("p:cSld")).get("name") == ADDED_SLIDE_NAME]
