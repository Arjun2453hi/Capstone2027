"""preview.py — a rough picture of a .pptx slide, drawn from its shapes.

For reviewing a repaired slide in the demo without PowerPoint: real
pictures at their real positions and sizes, text wrapped at its own
font size, other shapes as outlines. It shows the layout -- what sits
where, whether text and pictures collide -- not exact fonts or styling.
"""
from __future__ import annotations

import io
import textwrap
from typing import Callable, Tuple

from PIL import Image, ImageDraw, ImageFont
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

Rect = Tuple[float, float, float, float]
_FONT_PATHS = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",  # macOS
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # most Linux
    "C:/Windows/Fonts/arial.ttf",
)
_TEXT, _OUTLINE, _ACCENT = "#1d2124", "#b8c2cc", "#c55a11"


def _font(px: int):
    for path in _FONT_PATHS:
        try:
            return ImageFont.truetype(path, max(px, 6))
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=max(px, 6))
    except TypeError:  # Pillow < 10.1
        return ImageFont.load_default()


def _group_transform(group, outer: Callable[[Rect], Rect]) -> Callable[[Rect], Rect]:
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))

    def pair(tag, keys):
        node = xfrm.find(qn(tag))
        return tuple(int(node.get(k)) for k in keys) if node is not None else None

    off, ext = pair("a:off", ("x", "y")), pair("a:ext", ("cx", "cy"))
    ch_off, ch_ext = pair("a:chOff", ("x", "y")) or off, pair("a:chExt", ("cx", "cy")) or ext
    sx, sy = ext[0] / max(ch_ext[0], 1), ext[1] / max(ch_ext[1], 1)
    return lambda r: outer((
        off[0] + (r[0] - ch_off[0]) * sx,
        off[1] + (r[1] - ch_off[1]) * sy,
        off[0] + (r[2] - ch_off[0]) * sx,
        off[1] + (r[3] - ch_off[1]) * sy,
    ))


def render_slide(slide, slide_width: int, slide_height: int, width_px: int = 1000) -> Image.Image:
    scale = width_px / slide_width
    image = Image.new("RGB", (width_px, int(slide_height * scale)), "white")
    draw = ImageDraw.Draw(image)

    def px(r: Rect):
        return [int(v * scale) for v in r]

    def visit(shapes, to_slide: Callable[[Rect], Rect]):
        for shape in shapes:
            rect = to_slide((shape.left, shape.top, shape.left + shape.width, shape.top + shape.height))
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, _group_transform(shape, to_slide))
                continue
            b = px(rect)
            if b[2] <= b[0] or b[3] <= b[1]:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    picture = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
                    image.paste(picture.resize((b[2] - b[0], b[3] - b[1])), (b[0], b[1]))
                except Exception:  # an image format PIL cannot read (e.g. EMF)
                    draw.rectangle(b, outline=_ACCENT)
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
                font = _font(int((sizes[0] if sizes else 18) * 12700 * scale))
                height = getattr(font, "size", 12)
                y = b[1]
                for paragraph in shape.text_frame.paragraphs:
                    # Paragraph margins narrow the line, as PowerPoint does
                    # (slide 137 wraps its bullets before a diagram this way).
                    p_pr = paragraph._p.find(qn("a:pPr"))
                    margins = sum(int(p_pr.get(k, 0)) for k in ("marL", "marR")) if p_pr is not None else 0
                    per_line = max(int(((b[2] - b[0]) - margins * scale) / (height * 0.5)), 1)
                    # "\v" is a manual line break inside a paragraph (python-pptx's
                    # paragraph.text); the slide's own wrapping, kept as it is.
                    for piece in paragraph.text.split("\v"):
                        for line in textwrap.wrap(piece, per_line) or [""]:
                            draw.text((b[0], y), line, fill=_TEXT, font=font)
                            y += int(height * 1.2)
                continue
            draw.rectangle(b, outline=_OUTLINE)

    visit(slide.shapes, lambda r: r)
    return image
