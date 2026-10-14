"""Reconstruct a PDF as PowerPoint text boxes over raster diagram backgrounds.

This is a prototype conversion, not recovery of original PowerPoint objects.
Text spans become editable; diagrams and image-only pages remain raster images.
"""
import argparse
from io import BytesIO
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt


def convert(source: Path, output: Path):
    with pymupdf.open(source) as pdf:
        presentation = Presentation()
        presentation.slide_width = Pt(pdf[0].rect.width)
        presentation.slide_height = Pt(pdf[0].rect.height)
        text_boxes = 0
        for index, page in enumerate(pdf):
            width, height = page.rect.width, page.rect.height
            sx = pdf[0].rect.width / width
            sy = pdf[0].rect.height / height
            spans = [span for block in page.get_text("dict")["blocks"]
                     if block["type"] == 0 for line in block["lines"]
                     for span in line["spans"] if span["text"].strip()]
            # Keep diagrams/images in the background, remove only PDF text.
            for span in spans:
                page.add_redact_annot(pymupdf.Rect(span["bbox"]), fill=False)
            page.apply_redactions(images=0, graphics=0, text=0)
            background = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(BytesIO(background.tobytes("png")), 0, 0,
                                     presentation.slide_width, presentation.slide_height)
            for span in spans:
                x0, y0, x1, y1 = span["bbox"]
                shape = slide.shapes.add_textbox(Pt(x0 * sx), Pt(y0 * sy),
                                                Pt(max(x1 - x0 + 3, 1) * sx),
                                                Pt(max(y1 - y0 + 3, 1) * sy))
                frame = shape.text_frame
                frame.margin_left = frame.margin_right = 0
                frame.margin_top = frame.margin_bottom = 0
                frame.word_wrap = False
                frame.vertical_anchor = MSO_ANCHOR.TOP
                paragraph = frame.paragraphs[0]
                paragraph.space_before = paragraph.space_after = Pt(0)
                run = paragraph.add_run()
                run.text = span["text"]
                run.font.name = span["font"].split("+")[-1]
                run.font.size = Pt(span["size"] * sy)
                run.font.bold = bool(span["flags"] & 16)
                run.font.italic = bool(span["flags"] & 2)
                color = span["color"]
                run.font.color.rgb = RGBColor((color >> 16) & 255,
                                             (color >> 8) & 255, color & 255)
                shape.name = f"pdf_slide_{index}_text_{text_boxes}"
                text_boxes += 1
            slide.notes_slide.notes_text_frame.text = f"Source PDF slide_id={index}; slide_number={index + 1}"
            if (index + 1) % 25 == 0:
                print(f"Converted {index + 1}/{len(pdf)} slides", flush=True)
        output.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(output)
        saved = Presentation(output)
        if len(saved.slides) != len(pdf):
            raise RuntimeError("Converted slide count does not match PDF")
        print(f"Saved {output}: {len(saved.slides)} slides, {text_boxes} editable text boxes")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path(__file__).resolve().parents[1] / "data/cc-unit2-slides.pdf")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "data/cc-unit2-slides.pptx")
    args = parser.parse_args()
    convert(args.input, args.output)
