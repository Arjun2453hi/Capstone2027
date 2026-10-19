"""Small fake deck / dossier / report / .pptx builders shared by this
stage's tests. Nothing here touches a real model or the real deck."""
from __future__ import annotations

from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.util import Inches

from ...src._upstream import Deck, GapReport, QuestionMatch, Slide, TopicDossier
from ...src.planner import EditGenerator
from ...src.schema import SlideEdit

QUESTIONS = ["What is a cgroup?", "Can you show a cgroup hierarchy example?"]


def make_deck(texts: List[str]) -> Deck:
    slides = [
        Slide(slide_id=i, slide_number=i + 1, title=t.splitlines()[0] if t else None, title_font_size=18.0, raw_text=t)
        for i, t in enumerate(texts)
    ]
    return Deck(source_pdf="fake.pdf", num_pages=len(slides), slides=slides)


def make_dossier(topic_id: int = 0, slide_ids=(0, 1, 2), questions=QUESTIONS) -> TopicDossier:
    matches = [QuestionMatch(question=q, score=0.8, method="semantic", best_slide_id=slide_ids[0]) for q in questions]
    return TopicDossier(
        topic_id=topic_id, slide_ids=list(slide_ids), slide_range="", boundary_confidence=0.0,
        window_text="", cluster_size=len(matches), matched_questions=matches,
    )


def make_report(topic_id: int = 0, gap_type: str = "shallow_coverage", confidence: float = 0.85) -> GapReport:
    return GapReport(topic_id=topic_id, gap_type=gap_type, confidence=confidence, report_text="No hierarchy example is given.")


def make_edit(**overrides) -> SlideEdit:
    fields = dict(
        topic_id=0,
        after_slide_id=1,
        title="A cgroup hierarchy, step by step",
        bullets=[
            "The root cgroup /sys/fs/cgroup holds every process until one is moved.",
            "Creating /myapp under it gives a child group with its own limit files.",
            "Writing 50000 to cpu.max of /myapp caps the whole subtree at half a core.",
        ],
        addresses_questions=["Can you show a cgroup hierarchy example?"],
        gap_type="shallow_coverage",
        report_confidence=0.85,
    )
    fields.update(overrides)
    return SlideEdit(**fields)


class ScriptedGenerator(EditGenerator):
    def __init__(self, response):
        self.response = response
        self.calls = []

    def generate(self, system: str, user: str) -> str:
        self.calls.append((system, user))
        return self.response


def make_pptx(path: Path, titles: List[str]) -> Path:
    """A deck in the real one's shape: a title placeholder holding a
    banner line and the slide title, a thin rule under it, and a body
    text box."""
    prs = Presentation()
    layout = prs.slide_layouts[5]  # Title Only
    for title in titles:
        slide = prs.slides.add_slide(layout)
        frame = slide.shapes.title.text_frame
        frame.text = "COURSE BANNER"
        frame.add_paragraph().add_run().text = title
        rule = slide.shapes.build_freeform(Inches(0), Inches(1.4)).add_line_segments([(Inches(6), Inches(1.4))], close=False).convert_to_shape()
        rule.height = Inches(0.02)
        slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(8), Inches(3)).text_frame.text = f"body of {title}"
    prs.save(str(path))
    return path


def slide_titles(path: Path) -> List[str]:
    return [s.shapes.title.text_frame.paragraphs[-1].text for s in Presentation(str(path)).slides]
