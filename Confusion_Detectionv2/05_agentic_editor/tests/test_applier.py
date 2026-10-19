"""applier.py tests against a small generated .pptx: where new slides
land, that verbatim repeats of the anchor get them too, that the source
file is untouched, and that the new slide carries the deck's header."""
from __future__ import annotations

import hashlib

import pytest
from pptx import Presentation

from ..src.applier import BODY_SHAPE_NAME, added_slide_positions, apply_edits, duplicate_anchor_ids
from ..src.storage import load_plan_json, save_change_log_json, save_plan_json
from ..src.schema import CheckedEdit
from .fixtures.sample_inputs import make_deck, make_edit, make_pptx, slide_titles

TITLES = ["Intro", "Hierarchy", "Limits", "Summary"]


def _setup(tmp_path, titles=TITLES):
    source = make_pptx(tmp_path / "deck.pptx", titles)
    deck = make_deck([f"{t}\nbody of {t}" for t in titles])
    return source, deck


def test_new_slide_lands_directly_after_its_anchor_and_source_is_untouched(tmp_path):
    source, deck = _setup(tmp_path)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "out" / "edited.pptx"

    applied = apply_edits(source, output, [make_edit(after_slide_id=1, title="Worked hierarchy")], deck)

    assert slide_titles(output) == ["Intro", "Hierarchy", "Worked hierarchy", "Limits", "Summary"]
    assert applied[0].new_slide_positions == [2]
    assert added_slide_positions(output) == [2]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert slide_titles(source) == TITLES


def test_new_slide_carries_the_decks_header_and_its_own_bullets(tmp_path):
    source, deck = _setup(tmp_path)
    edit = make_edit(after_slide_id=1, title="Worked hierarchy")
    output = tmp_path / "edited.pptx"
    apply_edits(source, output, [edit], deck)

    slide = Presentation(str(output)).slides[2]
    title_lines = [p.text for p in slide.shapes.title.text_frame.paragraphs]
    assert title_lines == ["COURSE BANNER", "Worked hierarchy"]  # banner kept, only the title line replaced
    body = next(s for s in slide.shapes if s.name == BODY_SHAPE_NAME)
    assert [p.text for p in body.text_frame.paragraphs] == edit.bullets
    assert len([s for s in slide.shapes if s.shape_type == 5]) == 1  # the rule under the header
    assert edit.addresses_questions[0] in slide.notes_slide.notes_text_frame.text


def test_several_edits_keep_plan_order_and_report_final_positions(tmp_path):
    source, deck = _setup(tmp_path)
    edits = [
        make_edit(after_slide_id=2, title="After limits"),
        make_edit(after_slide_id=0, title="After intro A"),
        make_edit(after_slide_id=0, title="After intro B"),
    ]
    output = tmp_path / "edited.pptx"

    applied = apply_edits(source, output, edits, deck)

    titles = slide_titles(output)
    assert titles == ["Intro", "After intro A", "After intro B", "Hierarchy", "Limits", "After limits", "Summary"]
    for a in applied:  # every reported position really holds that edit's slide
        assert [titles[p] for p in a.new_slide_positions] == [a.edit.title]


def test_edit_is_mirrored_into_a_verbatim_repeat_of_the_lecture_but_not_a_recurring_slide(tmp_path):
    """Slides 1-2 repeat as 4-5 (a lecture included twice). "Thanks"
    also recurs, but what precedes it differs, so it is not a repeat."""
    titles = ["Intro", "Hierarchy", "Limits", "Thanks", "Hierarchy", "Limits", "Other", "Thanks"]
    source = make_pptx(tmp_path / "deck.pptx", titles)
    deck = make_deck(["Intro\nx", "Hierarchy\nsame", "Limits\nsame", "Thanks", "Hierarchy\nsame", "Limits\nsame", "Other\ny", "Thanks"])
    assert duplicate_anchor_ids(deck, 2) == [5]
    assert duplicate_anchor_ids(deck, 3) == []  # "Thanks" at 7 follows a different slide
    output = tmp_path / "edited.pptx"

    applied = apply_edits(source, output, [make_edit(after_slide_id=2, title="New")], deck)

    assert slide_titles(output) == ["Intro", "Hierarchy", "Limits", "New", "Thanks", "Hierarchy", "Limits", "New", "Other", "Thanks"]
    assert applied[0].new_slide_positions == [3, 7]
    assert applied[0].mirrored_after_slide_ids == [5]


def test_refuses_to_overwrite_the_source_or_work_on_a_misaligned_deck(tmp_path):
    source, deck = _setup(tmp_path)
    with pytest.raises(ValueError, match="never modified"):
        apply_edits(source, source, [make_edit()], deck)
    with pytest.raises(ValueError, match="would not line up"):
        apply_edits(source, tmp_path / "o.pptx", [make_edit()], make_deck(["only one slide"]))


def test_plan_and_change_log_round_trip(tmp_path):
    source, deck = _setup(tmp_path)
    checked = [CheckedEdit(make_edit()), CheckedEdit(make_edit(title="bad"), problems=["title is empty"])]
    save_plan_json(checked, {4: "covered: nothing to add"}, tmp_path / "plan.json")

    loaded = load_plan_json(tmp_path / "plan.json")
    assert [c.edit for c in loaded] == [c.edit for c in checked]
    assert [c.accepted for c in loaded] == [True, False]

    applied = apply_edits(source, tmp_path / "edited.pptx", [loaded[0].edit], deck)
    save_change_log_json(applied, source, tmp_path / "edited.pptx", tmp_path / "log.json")
    import json

    log = json.loads((tmp_path / "log.json").read_text())
    assert log["slides_added"] == 1 and log["edits"][0]["new_slide_numbers"] == [3]


def test_header_is_borrowed_from_the_nearest_slide_that_has_the_standard_one(tmp_path):
    """The anchor and its neighbour have a title placeholder but not the
    deck's header (one line, no banner, no rule) -- as some slides of
    the real deck do. The new slide must take its header from the
    standard slide further away, not copy the odd one."""
    source, deck = _setup(tmp_path)
    prs = Presentation(str(source))
    for index in (1, 2):  # strip "Hierarchy" and "Limits" down to a bare one-line title
        slide = prs.slides[index]
        frame = slide.shapes.title.text_frame
        frame.text = f"odd slide {index}"
        for shape in [s for s in slide.shapes if s.shape_type == 5]:
            shape._element.getparent().remove(shape._element)
    prs.save(str(source))
    output = tmp_path / "edited.pptx"

    apply_edits(source, output, [make_edit(after_slide_id=1, title="Worked hierarchy")], deck)

    new_slide = Presentation(str(output)).slides[2]
    assert [p.text for p in new_slide.shapes.title.text_frame.paragraphs] == ["COURSE BANNER", "Worked hierarchy"]
    assert len([s for s in new_slide.shapes if s.shape_type == 5]) == 1
