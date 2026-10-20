"""Tests for adding bullets to existing slides and repairing broken ones:
slide_parts.py (what a repair carries over), repair.py (the planner and
its one retry), the checker's rules for both operations, the applier,
and the preview drawing."""
from __future__ import annotations

import io
import json

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..src.applier import BODY_SHAPE_NAME, apply_edits, fit_font_pt
from ..src.checker import MAX_APPEND_TOTAL_CHARS, check_edit, check_edits
from ..src.preview import render_slide
from ..src.repair import RepairPlanner, parse_repair
from ..src.schema import APPEND_BULLETS, REPAIR_SLIDE, REPAIR_TOPIC_ID, SlideEdit
from ..src.slide_parts import deck_banner, split_slide
from .fixtures.sample_inputs import ScriptedGenerator, make_deck, make_dossier, make_edit, make_pptx, slide_titles

TITLES = ["Intro", "Hierarchy", "Limits", "Summary"]
PARAGRAPH = "Containers share the host kernel and isolate processes with namespaces and cgroups."


def _png(size=(40, 40), colour="gray") -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _broken_deck(path):
    """make_pptx's deck, with slide 2 ("Limits") broken the way the real
    deck's are: a paragraph under a stretched picture, a label on the
    picture, two labels printed over each other, and a slide-sized
    backdrop picture."""
    make_pptx(path, TITLES)
    prs = Presentation(str(path))
    slide = prs.slides[2]
    slide.shapes.add_picture(_png((10, 10), "white"), 0, 0, prs.slide_width, prs.slide_height)  # backdrop
    slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(6), Inches(1.5)).text_frame.text = PARAGRAPH
    slide.shapes.add_picture(_png((400, 200)), Inches(1), Inches(2), Inches(2), Inches(3))  # 2:1 image shown 2:3
    slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(1), Inches(0.4)).text_frame.text = "Kernel"
    for text in ("Ring 3", "Ring 0"):  # stacked on each other
        slide.shapes.add_textbox(Inches(1.2), Inches(4), Inches(1), Inches(0.4)).text_frame.text = text
    prs.save(str(path))
    return path


def _parts(path, index=2):
    prs = Presentation(str(path))
    return split_slide(prs.slides[index], prs.slide_width, prs.slide_height, deck_banner(prs))


# ---- slide_parts ---------------------------------------------------------


def test_split_keeps_paragraphs_for_the_llm_and_moves_pictures_and_their_labels(tmp_path):
    parts = _parts(_broken_deck(tmp_path / "d.pptx"))

    assert parts.title == "Limits"  # the "COURSE BANNER" line is not part of it
    assert parts.text_blocks == ["body of Limits", PARAGRAPH]
    pictures = [g for g in parts.graphics if not g.is_label]
    labels = [g for g in parts.graphics if g.is_label]
    assert len(pictures) == 1  # the backdrop is dropped
    assert len(labels) == 1  # "Kernel"; the stacked "Ring 3"/"Ring 0" are dropped
    left, top, right, bottom = pictures[0].rect
    assert abs((right - left) / (bottom - top) - 2.0) < 0.01  # true 2:1 proportions restored


def test_the_least_stretched_copy_of_a_repeated_picture_is_kept(tmp_path):
    path = make_pptx(tmp_path / "d.pptx", TITLES)
    prs = Presentation(str(path))
    slide = prs.slides[1]
    slide.shapes.add_picture(_png((300, 100)), Inches(0.5), Inches(2), Inches(6), Inches(1))  # 6:1, image is 3:1
    slide.shapes.add_picture(_png((300, 100)), Inches(0.5), Inches(4), Inches(3.3), Inches(1))  # 3.3:1
    prs.save(str(path))

    pictures = [g for g in _parts(path, 1).graphics if not g.is_label]
    assert len(pictures) == 1
    assert pictures[0].rect[1] > Inches(3.5)  # the second, nearly undistorted copy


# ---- repair planner ------------------------------------------------------


def _repair_response(title="Limits", bullets=(PARAGRAPH, "body of Limits")):
    return json.dumps({"title": title, "bullets": list(bullets)})


def test_repair_planner_sends_only_the_slides_own_text_and_returns_a_repair_edit(tmp_path):
    parts = _parts(_broken_deck(tmp_path / "d.pptx"))
    generator = ScriptedGenerator(_repair_response())

    edit = RepairPlanner(generator).plan_repair(parts, slide_id=2)

    assert (edit.operation, edit.topic_id, edit.after_slide_id) == (REPAIR_SLIDE, REPAIR_TOPIC_ID, 2)
    assert check_edit(edit, None, None).accepted
    assert len(generator.calls) == 1  # accepted first time: no retry
    _, user = generator.calls[0]
    assert PARAGRAPH in user and "Diagram labels" in user and "Kernel" in user


def test_a_rejected_repair_is_sent_back_once_with_the_checkers_findings(tmp_path):
    parts = _parts(_broken_deck(tmp_path / "d.pptx"))
    invented = _repair_response(bullets=("Kubernetes schedules pods across worker nodes automatically.", "Etcd stores cluster state durably."))

    class TwoAnswers(ScriptedGenerator):
        def generate(self, system, user):
            self.calls.append((system, user))
            return invented if len(self.calls) == 1 else _repair_response()

    generator = TwoAnswers(None)
    edit = RepairPlanner(generator).plan_repair(parts, slide_id=2)

    assert len(generator.calls) == 2
    assert "Your previous rebuild was rejected" in generator.calls[1][1]
    assert "not on the original slide" in generator.calls[1][1]
    assert check_edit(edit, None, None).accepted


def test_unusable_repair_output_yields_none():
    for raw in [None, "", "not json", "[1]", '{"title": "x"}']:
        assert parse_repair(raw, 0, "source") is None


# ---- checker -------------------------------------------------------------


def _repair(bullets, source="Containers share the host kernel and isolate processes with namespaces and cgroups. Limits"):
    return SlideEdit(REPAIR_TOPIC_ID, 2, "Limits", list(bullets), [], operation=REPAIR_SLIDE, source_text=source)


def test_a_repair_may_not_add_content_or_drop_most_of_the_slide():
    faithful = _repair(["Containers share the host kernel.", "They isolate processes with namespaces and cgroups."])
    assert check_edit(faithful, None, None).accepted

    padded = _repair(["Containers share the host kernel.", "Kubernetes orchestrates deployments across clusters automatically."])
    assert any("not on the original slide" in p for p in check_edit(padded, None, None).problems)

    thin = _repair(["Containers share things.", "Limits apply here."])
    assert any("keeps only" in p for p in check_edit(thin, None, None).problems)


def test_appending_is_limited_to_a_few_bullets_on_a_slide_with_room():
    dossier = make_dossier()
    roomy = make_deck(["Cgroups\nlimits", "Hierarchy\ncgroups are hierarchical", "Limits\nmemory"])
    full = make_deck(["Cgroups\nlimits", "Hierarchy\n" + "x " * (MAX_APPEND_TOTAL_CHARS // 2), "Limits\nmemory"])
    edit = make_edit(operation=APPEND_BULLETS, title="", bullets=make_edit().bullets[:2])

    assert check_edit(edit, dossier, roomy).accepted
    assert any("would hold" in p and "overflow" in p for p in check_edit(edit, dossier, full).problems)
    four = make_edit(operation=APPEND_BULLETS, title="", bullets=make_edit().bullets * 2)
    assert any("bullets (need 1-3)" in p for p in check_edit(four, dossier, roomy).problems)


def test_repairs_are_not_compared_with_gap_edits_for_repeats():
    deck = make_deck(["a", "b", "c"])
    checked = check_edits([make_edit(after_slide_id=2), _repair(["Containers share the host kernel.", "They isolate processes with namespaces and cgroups."])], {0: make_dossier()}, deck)
    assert all(c.accepted for c in checked)


# ---- applier -------------------------------------------------------------


def test_appended_bullets_copy_the_last_bullets_formatting_and_shift_nothing(tmp_path):
    source = make_pptx(tmp_path / "deck.pptx", TITLES)
    prs = Presentation(str(source))
    body = next(s for s in prs.slides[1].shapes if s.has_text_frame and s.text_frame.text == "body of Hierarchy")
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    prs.save(str(source))
    deck = make_deck([f"{t}\nbody of {t}" for t in TITLES])
    edits = [make_edit(after_slide_id=0, title="New after intro"), make_edit(operation=APPEND_BULLETS, after_slide_id=1, bullets=["First added point here.", "Second added point here."])]

    applied = apply_edits(source, tmp_path / "out.pptx", edits, deck)

    assert slide_titles(tmp_path / "out.pptx") == ["Intro", "New after intro", "Hierarchy", "Limits", "Summary"]
    assert applied[1].new_slide_positions == [2]  # Hierarchy, shifted by the slide inserted before it
    slide = Presentation(str(tmp_path / "out.pptx")).slides[2]
    body = next(s for s in slide.shapes if s.has_text_frame and "body of Hierarchy" in s.text_frame.text)
    paragraphs = body.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["body of Hierarchy", "First added point here.", "Second added point here."]
    assert all(p.runs[0].font.size == Pt(20) for p in paragraphs)
    assert "added the last 2 bullet(s)" in slide.notes_slide.notes_text_frame.text


def test_repair_inserts_the_rebuild_after_the_original_and_hides_the_original(tmp_path):
    source = _broken_deck(tmp_path / "deck.pptx")
    deck = make_deck([f"{t}\nbody of {t}" for t in TITLES])
    edit = _repair(["Containers share the host kernel.", "They isolate processes with namespaces and cgroups."])

    applied = apply_edits(source, tmp_path / "out.pptx", [edit], deck)

    prs = Presentation(str(tmp_path / "out.pptx"))
    assert len(prs.slides) == 5
    assert applied[0].hidden_slide_positions == [2] and applied[0].new_slide_positions == [3]
    assert prs.slides[2]._element.get("show") == "0"  # the broken original: kept, hidden
    rebuilt = prs.slides[3]
    assert rebuilt._element.get("show") is None
    assert [p.text for p in rebuilt.shapes.title.text_frame.paragraphs] == ["COURSE BANNER", "Limits"]
    body = next(s for s in rebuilt.shapes if s.name == BODY_SHAPE_NAME)
    assert [p.text for p in body.text_frame.paragraphs] == edit.bullets
    pictures = [s for s in rebuilt.shapes if s.shape_type == 13]
    assert len(pictures) == 1 and pictures[0].image.size == (400, 200)  # the image's relationship was re-linked
    assert any(s.has_text_frame and s.text_frame.text == "Kernel" for s in rebuilt.shapes)
    ids = [n.get("id") for n in rebuilt.shapes._spTree.iter(qn("p:cNvPr"))]
    assert len(ids) == len(set(ids))
    left_of_picture = body.left + body.width <= pictures[0].left or body.top + body.height <= pictures[0].top
    assert left_of_picture  # text and picture do not overlap


def test_font_fitting_shrinks_long_text_but_not_below_the_minimum():
    one_line = fit_font_pt(["short"], Inches(6), Inches(4))
    crowded = fit_font_pt(["word " * 60] * 6, Inches(6), Inches(4))
    assert one_line == 22 and 12 <= crowded < 22


def test_preview_draws_a_slide_with_its_pictures(tmp_path):
    source = _broken_deck(tmp_path / "deck.pptx")
    prs = Presentation(str(source))
    image = render_slide(prs.slides[2], prs.slide_width, prs.slide_height, width_px=500)
    assert image.size == (500, int(500 * prs.slide_height / prs.slide_width))
    assert len(set(image.getdata())) > 2  # not blank
