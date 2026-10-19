"""layout_check.py tests: each signal fires on a deck built to break it,
and a clean slide is not flagged."""
from __future__ import annotations

import io

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from ..src.layout_check import find_layout_issues

PARAGRAPH = "Containers share the host kernel and isolate processes with namespaces and cgroups."


def _png() -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 40), "gray").save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _deck(path):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    clean = prs.slides.add_slide(blank)  # paragraph on the left, picture on the right
    clean.shapes.add_textbox(Inches(0.5), Inches(1), Inches(4), Inches(2)).text_frame.text = PARAGRAPH
    clean.shapes.add_picture(_png(), Inches(5.5), Inches(1), Inches(3), Inches(3))

    covered = prs.slides.add_slide(blank)  # picture drawn over the paragraph
    covered.shapes.add_textbox(Inches(0.5), Inches(1), Inches(6), Inches(2)).text_frame.text = PARAGRAPH
    covered.shapes.add_picture(_png(), Inches(1), Inches(1), Inches(5), Inches(2))

    stacked = prs.slides.add_slide(blank)  # two paragraphs on top of each other
    for _ in range(2):
        stacked.shapes.add_textbox(Inches(0.5), Inches(1), Inches(6), Inches(2)).text_frame.text = PARAGRAPH
    prs.save(str(path))
    return path


def _pdf(path):
    c = canvas.Canvas(str(path), pagesize=(720, 405))
    c.drawString(40, 300, PARAGRAPH)  # page 0: one clean line
    c.showPage()
    c.drawString(40, 300, PARAGRAPH)  # page 1: two build steps printed at once
    c.drawString(43, 302, "Hardware traps to the VMM which vectors control back to the guest OS")
    c.showPage()
    c.drawString(40, 300, PARAGRAPH)
    c.showPage()
    c.save()
    return path


def test_geometry_flags_text_under_a_picture_and_text_on_text_but_not_a_clean_slide(tmp_path):
    issues = {i.slide_id: i for i in find_layout_issues(_deck(tmp_path / "d.pptx"))}

    assert 0 not in issues
    assert any("picture or diagram overlaps" in r for r in issues[1].reasons)
    assert any("two paragraphs overlap" in r for r in issues[2].reasons)


def test_rendering_flags_characters_drawn_over_other_characters(tmp_path):
    issues = {i.slide_id: i for i in find_layout_issues(_deck(tmp_path / "d.pptx"), _pdf(tmp_path / "d.pdf"))}

    assert any("drawn on top of other text" in r for r in issues[1].reasons)
    assert not any("drawn on top" in r for i in issues.values() if i.slide_id != 1 for r in i.reasons)
    assert issues[1].score > issues[2].score  # rendered evidence outranks geometry alone
