"""app.py — live demo of the gap-to-slide-edit loop.

    streamlit run 05_agentic_editor/app.py

from the Confusion_Detectionv2 directory. Pick a gap found by stage 4,
see its report, let the LLM suggest slides (checked and fact-checked),
approve the ones you want, and apply them to a copy of the deck.
"""
from __future__ import annotations

import copy
import html
import importlib
import json
import os
import subprocess
import sys
from pathlib import Path

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import streamlit as st  # noqa: E402
from dotenv import load_dotenv  # noqa: E402

load_dotenv(_PROJECT_ROOT / ".env")

editor = importlib.import_module("05_agentic_editor.src")
applier = importlib.import_module("05_agentic_editor.src.applier")
editor_storage = importlib.import_module("05_agentic_editor.src.storage")
upstream = importlib.import_module("05_agentic_editor.src._upstream")

DECK_JSON = _PROJECT_ROOT / "01_deck_parsing" / "parsed_deck.json"
GAP_INPUT = _PROJECT_ROOT / "03_question_mapping" / "gap_verification_input.json"
REPORTS_DIR = _PROJECT_ROOT / "04_gap_reporting_agent" / "output"
SOURCE_PDF = _PROJECT_ROOT / "data" / "cc-unit2-slides.pdf"
SOURCE_PPTX = _PROJECT_ROOT / "data" / "cc-unit2-slides.pptx"
OUTPUT_DIR = Path(__file__).resolve().parent / "output"
DEMO_PPTX = OUTPUT_DIR / "cc-unit2-slides.demo.pptx"
DEMO_LOG = OUTPUT_DIR / "demo_change_log.json"
layout_check = importlib.import_module("05_agentic_editor.src.layout_check")
repair = importlib.import_module("05_agentic_editor.src.repair")
preview = importlib.import_module("05_agentic_editor.src.preview")
schema = importlib.import_module("05_agentic_editor.src.schema")
revise = importlib.import_module("05_agentic_editor.src.revise")
planner_module = importlib.import_module("05_agentic_editor.src.planner")
REVIEWED = OUTPUT_DIR / "reviewed_suggestions.json"
FACT_CHECK_MAX_TOKENS = 1500  # the fact-check's answer is short: a smaller budget waits less on the rate limit
REPAIRS = schema.REPAIR_TOPIC_ID  # key of the repairs among state.suggestions

GAP_LABELS = {
    "complete_omission": "Missing entirely",
    "shallow_coverage": "Covered too thinly",
    "fragmented_context": "Scattered across slides",
    "covered": "Covered",
}
GAP_COLOURS = {"complete_omission": "red", "shallow_coverage": "orange", "fragmented_context": "violet", "covered": "green"}
GAP_MEANINGS = {
    "complete_omission": "The deck never explains what students asked about.",
    "shallow_coverage": "The deck mentions it, but too briefly to answer the students' questions.",
    "fragmented_context": "The explanation exists, but is split across slides far apart.",
    "covered": "The deck already answers these questions.",
}

st.set_page_config(page_title="Slide Gap Editor", layout="wide")


@st.cache_resource
def load_inputs():
    deck = upstream.load_deck_json(DECK_JSON)
    gap_input = upstream.load_gap_verification_input_json(GAP_INPUT)
    dossiers = {d.topic_id: d for d in gap_input.topics}
    unmatched = [u.question for u in gap_input.unmatched_questions]
    reports = {}
    for path in sorted(REPORTS_DIR.glob("topic_*_report.json")):
        report = upstream.GapReport(**json.loads(path.read_text(encoding="utf-8")))
        reports[report.topic_id] = report
    index = json.loads((REPORTS_DIR / "index.json").read_text(encoding="utf-8"))
    return deck, dossiers, reports, index, unmatched


@st.cache_resource
def room_checker():
    return applier.RoomChecker(SOURCE_PPTX)


@st.cache_resource
def llm(api_key):
    generator = editor.GroqEditGenerator(api_key=api_key)
    verifier = editor.EditVerifier(planner_module.with_max_tokens(generator, FACT_CHECK_MAX_TOKENS))
    return editor.EditPlanner(generator), verifier, generator.model, repair.RepairPlanner(generator)


@st.cache_data(show_spinner="Drawing the changed slide…")
def _changed_slide_image(edit_json: str):
    """The slide an edit produces (the new slide, the slide with added
    bullets, or the rebuilt slide), drawn after really applying the edit
    to a scratch copy of the deck -- what the board sees is what the
    file will contain."""
    import tempfile

    from pptx import Presentation

    edit = schema.SlideEdit(**json.loads(edit_json))
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "preview.pptx"
        applied = applier.apply_edits(SOURCE_PPTX, out, [edit], deck)
        prs = Presentation(str(out))
        return preview.render_slide(prs.slides[applied[0].new_slide_positions[0]], prs.slide_width, prs.slide_height, width_px=900)


def changed_slide_image(edit):
    from dataclasses import asdict

    return _changed_slide_image(json.dumps(asdict(edit), sort_keys=True))


rebuilt_preview = changed_slide_image  # repairs use the same drawing


@st.cache_data
def slide_image(slide_id: int):
    import pypdfium2

    pdf = pypdfium2.PdfDocument(str(SOURCE_PDF))
    return pdf[slide_id].render(scale=0.8).to_pil()


@st.cache_data
def broken_layout_candidates():
    return layout_check.find_layout_issues(SOURCE_PPTX, SOURCE_PDF, max_results=10)


def edit_key(edit) -> str:
    return f"{edit.topic_id}|{edit.after_slide_id}|{edit.title}"


deck, dossiers, reports, index, unmatched = load_inputs()
reviewed_topics, reviewed_repairs, reviewed_meta = editor_storage.load_reviewed_json(REVIEWED)
state = st.session_state
state.setdefault("suggestions", {})  # topic_id -> [CheckedEdit]
state.setdefault("approved", {})  # edit_key -> bool
state.setdefault("applied", None)

st.title("Slide Gap Editor")
st.caption("Student questions → gaps in the deck → LLM-written slides → checked → applied to a copy of the deck")

# ---- Step 1: pick a gap -------------------------------------------------
def plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


# Every report the agent wrote, in three groups (index.json order, worst first).
groups = {
    "Gaps to fix": [r for r in index if r["gap_type"] != "covered" and r["topic_id"] in dossiers and dossiers[r["topic_id"]].slide_ids],
    "Already covered": [r for r in index if r["gap_type"] == "covered"],
    "Off-topic questions": [r for r in index if r["topic_id"] == -1],
}
with st.sidebar:
    st.header("1. Pick a topic")
    group = st.radio(
        "Show",
        list(groups),
        format_func=lambda g: f"{g} ({len(groups[g])})" if g != "Off-topic questions" else f"{g} ({len(unmatched)})",
        horizontal=False,
    )
    st.caption({
        "Gaps to fix": "Topics where the agent found the deck falls short, worst first.",
        "Already covered": "Topics the agent checked and found already answered: nothing to fix.",
        "Off-topic questions": "Student questions about other subjects (databases, CSS, Java...), so there is no slide to add them to.",
    }[group])
    rows = groups[group]
    labels = {
        r["topic_id"]: (
            f"Topic {r['topic_id']} · {GAP_LABELS.get(r['gap_type'], r['gap_type'])} · {plural(r['backed_by_questions'], 'question')}"
            if r["topic_id"] != -1 else f"Questions outside this deck · {plural(len(unmatched), 'question')}"
        )
        for r in rows
    }
    topic_id = st.radio("Topic", list(labels), format_func=labels.get, label_visibility="collapsed")
    st.divider()
    use_reviewed = st.toggle(
        "Use pre-checked results",
        value=bool(reviewed_topics),
        disabled=not reviewed_topics,
        help="Suggestions and repairs generated ahead of time by this same pipeline and read by a person. "
        "Instant, and no API key or rate limit involved. Switch off to generate live with the LLM.",
    )
    if use_reviewed:
        st.caption(f"Pre-checked results from {(reviewed_meta.get('generated_at') or '')[:10]}, {reviewed_meta.get('model')}.")
    else:
        st.caption("Live: every suggestion is generated, checked and fact-checked now (about 1–3 minutes each on a free Groq key).")
    st.divider()
    # .env is not in git, so a fresh clone has no key: ask for one here.
    # It stays in this browser session only and is never written to disk.
    api_key = os.getenv("GROQ_API_KEY") or st.text_input(
        "Groq API key", type="password", help="Needed only to suggest slides. Free at console.groq.com/keys."
    )
    st.divider()
    n_approved = sum(state.approved.values())
    st.metric("Slides approved so far", n_approved)

report, dossier = reports[topic_id], dossiers.get(topic_id)
questions = [m.question for m in dossier.matched_questions] if dossier is not None else unmatched
has_slides = dossier is not None and bool(dossier.slide_ids)

# ---- Step 2: the gap report --------------------------------------------
st.header(f"2. What the gap report found — topic {topic_id}")
# The gap type gets a full-width line: in a metric box the longer names
# ("Scattered across slides") were cut off.
if has_slides or report.gap_type == "covered":
    colour, label, meaning = GAP_COLOURS.get(report.gap_type, "gray"), GAP_LABELS.get(report.gap_type, report.gap_type), GAP_MEANINGS.get(report.gap_type, "")
else:  # the unmatched questions: off-topic, not a gap in this deck
    colour, label, meaning = "gray", "Not about this deck", "Questions on other subjects; the agent searched the whole deck and found nothing related."
st.markdown(f"**Gap type:** :{colour}-badge[{label}] &nbsp; {meaning}")
c1, c2, c3 = st.columns(3)
c1.metric("Agent confidence", f"{report.confidence:.0%}")
c2.metric("Student questions", len(questions))
c3.metric("Slides in topic", f"{min(dossier.slide_ids) + 1}–{max(dossier.slide_ids) + 1}" if has_slides else "none")
left, right = st.columns([3, 2])
with left:
    with st.expander("Full report from the investigating agent", expanded=True):
        st.markdown(report.report_text)
with right:
    st.markdown("**Questions students asked about this topic**")
    for q in questions:
        st.markdown(f"- {q}")

reason = editor.skip_reason(report, dossier)
if report.gap_type == "covered":
    st.success("Nothing to fix: the agent checked these questions against the deck and found them already answered.")
elif not has_slides:
    st.info("These questions are about other subjects, not this deck, so there is no place in it to add a slide. "
            "The agent searched the whole deck to confirm that before concluding.")
elif reason:
    st.info(f"The editor does not act on this report: {reason}.")

# ---- Step 3: suggestions ------------------------------------------------
if reason is None:
    st.header("3. Changes the LLM suggests")
cached = reviewed_topics.get(topic_id) if use_reviewed else None
if reason is not None:
    pass  # nothing to suggest for this report; the note above says why
elif use_reviewed and cached is None:
    st.info("No pre-checked results for this gap; switch off “Use pre-checked results” in the sidebar to generate live.")
elif not use_reviewed and not api_key:
    st.info("Enter a Groq API key in the sidebar to have the LLM suggest slides.")
can_suggest = reason is None and (cached is not None if use_reviewed else bool(api_key))
if reason is None and st.button("Suggest slides for this gap", type="primary", disabled=not can_suggest):
    if use_reviewed:
        checked = copy.deepcopy(cached)
    else:
        try:
            planner, verifier, model, _ = llm(api_key)
        except Exception as e:  # missing/invalid key, or the model was retired
            st.error(f"Could not connect to Groq: {e}")
            st.stop()
        with st.status(f"Working with {model}…", expanded=True) as status:
            # Checked together with the slides already suggested for other
            # gaps, so a repeat of one of them (e.g. both copies of the
            # duplicated lecture asking for the same slide) is caught.
            earlier = [c.edit for t, cs in state.suggestions.items() if t not in (topic_id, REPAIRS) for c in cs if c.accepted]
            checked = revise.suggest_for_topic(
                report, dossier, deck, dossiers, planner, verifier, earlier=earlier,
                progress=lambda m: st.write(m + "…"), room=room_checker(),
            )
            error = getattr(planner.generator, "last_error", None)
            if not checked and error:
                status.update(label="The LLM call failed", state="error")
                st.error(f"Could not get suggestions from {model}: {error}. Wait a minute and try again.")
                st.stop()
            status.update(label=f"{sum(c.accepted for c in checked)} of {len(checked)} suggestion(s) passed every check", state="complete")
    state.suggestions[topic_id] = checked
    for c in checked:
        state.approved[edit_key(c.edit)] = c.accepted


def show_suggestion(number: int, c) -> None:
    """Top to bottom: what the change is, the exact text, whether it
    passed review, then the slide before and after."""
    e = c.edit
    where = e.after_slide_id + 1
    if e.operation == schema.APPEND_BULLETS:
        st.subheader(f"Suggestion {number}: add {plural(len(e.bullets), 'bullet')} to slide {where}")
    else:
        st.subheader(f"Suggestion {number}: new slide after slide {where}")

    st.markdown("**1. The change, in text**")
    with st.container(border=True):
        if e.operation == schema.APPEND_BULLETS:
            st.markdown(f"Added to the end of slide {where} (*{e.title or 'untitled'}*), in that slide's own bullet style:")
        else:
            st.markdown(f"**{e.title}**")
        st.markdown("\n".join(f"- {b}" for b in e.bullets))
        st.caption("Answers the students' questions: " + "; ".join(f"“{q}”" for q in e.addresses_questions))

    st.markdown("**2. Review**")
    if c.accepted:
        corrected = (
            f" The fact-check found a problem in the first draft, and the model corrected it "
            f"({plural(c.revisions, 'revision')}) before it passed."
            if c.revisions else ""
        )
        st.success("Passed the quality checks and the fact-check: consistent with the deck, and it answers the questions above." + corrected)
    else:
        st.caption("Why it was blocked: " + " · ".join(c.problems))

    st.markdown("**3. The slide, before and after**")
    before, after = st.columns(2)
    if e.operation == schema.APPEND_BULLETS:
        before.caption(f"Slide {where} now")
        after.caption(f"Slide {where} with the bullets added")
    else:
        before.caption(f"Existing slide {where} (unchanged)")
        after.caption("The new slide, right after it (drawn from the edited file; open the deck for exact fonts)")
    before.image(slide_image(e.after_slide_id), width="stretch")
    try:
        after.image(changed_slide_image(e), width="stretch")
    except Exception as err:  # the drawing is a convenience; never block the review on it
        after.caption(f"(No drawing available: {err})")


suggestions = state.suggestions.get(topic_id)
if suggestions is not None:
    passed = [c for c in suggestions if c.accepted]
    # Held back only because another gap's slide already fills this spot
    # (e.g. the second copy of the duplicated lecture): not a failure.
    covered = [c for c in suggestions if not c.accepted and all(p.startswith("repeats ") for p in c.problems)]
    blocked = [c for c in suggestions if not c.accepted and c not in covered]
    for c in covered:
        other = c.problems[0].split("(topic ")[1].split(")")[0] if "(topic " in c.problems[0] else "another"
        st.info(
            f"Already covered: the slide planned for topic {other} goes in this same place — in both copies of the "
            f"repeated lecture — so this gap needs no separate slide (its draft “{c.edit.title}” said the same thing)."
        )
    if not passed and not covered:
        st.info(
            "No suggestion for this gap passed every check, so none is offered. "
            "The blocked drafts and the reasons are below."
            if blocked else "The LLM proposed no slides for this gap."
        )
    for i, c in enumerate(passed, start=1):
        show_suggestion(i, c)
        key = edit_key(c.edit)
        state.approved[key] = st.checkbox("Approve this slide", value=state.approved.get(key, True), key=f"approve-{key}")
        st.divider()
    if blocked:
        with st.expander(f"Blocked by the safety checks ({len(blocked)}) — never applied unless you approve them"):
            for c in blocked:
                st.markdown(f"**{c.edit.title or 'Added bullets'}** (after slide {c.edit.after_slide_id + 1})")
                st.markdown("\n".join(f"- {b}" for b in c.edit.bullets))
                st.caption("Why it was blocked: " + " · ".join(c.problems))
                key = edit_key(c.edit)
                state.approved[key] = st.checkbox("Approve anyway", value=state.approved.get(key, False), key=f"approve-{key}")
                st.divider()

# ---- Step 4: apply -----------------------------------------------------
st.header("4. Apply the approved slides to the deck")
approved = [c.edit for checked in state.suggestions.values() for c in checked if state.approved.get(edit_key(c.edit))]
st.write(f"{len(approved)} slide(s) approved across all gaps. The original deck is never changed; a copy is written.")
if st.button("Apply approved slides", type="primary", disabled=not approved):
    applied = applier.apply_edits(SOURCE_PPTX, DEMO_PPTX, approved, deck)
    editor_storage.save_change_log_json(applied, SOURCE_PPTX, DEMO_PPTX, DEMO_LOG)
    state.applied = applied

if state.applied:
    total = sum(len(a.new_slide_positions) for a in state.applied if a.edit.operation != schema.APPEND_BULLETS or a.note)
    st.success(f"Done: {len(deck.slides)} slides → {len(deck.slides) + total} slides.")
    for a in state.applied:
        where = ", ".join(str(p + 1) for p in a.new_slide_positions)
        extra = " (also in the repeated copy of this lecture)" if a.mirrored_after_slide_ids else ""
        if a.edit.operation == schema.REPAIR_SLIDE:
            hidden = ", ".join(str(p + 1) for p in a.hidden_slide_positions)
            st.markdown(f"- **Repaired slide {a.edit.after_slide_id + 1}** → rebuilt as slide {where}; the original (slide {hidden}) is hidden, not deleted")
        elif a.edit.operation == schema.APPEND_BULLETS and not a.note:
            st.markdown(f"- **{len(a.edit.bullets)} bullet(s) added** to slide {where}{extra}")
        else:
            st.markdown(f"- **{a.edit.title}** → now slide {where}{extra}")
        if a.note:
            st.caption(a.note)
    d1, d2 = st.columns(2)
    d1.download_button("Download the edited deck", DEMO_PPTX.read_bytes(), file_name=DEMO_PPTX.name)
    if d2.button("Open it in PowerPoint"):
        subprocess.run(["open", "-a", "Microsoft PowerPoint", str(DEMO_PPTX)], check=False)

# ---- Broken layouts: detect, confirm, repair -----------------------------
st.divider()
st.header("Slides whose layout is broken")
st.caption(
    "Found from the slide geometry and the PDF rendering: text drawn over other text, pictures covering paragraphs. "
    "Some candidates are false alarms, so tick only the ones you can see are broken. A repair rebuilds the slide "
    "from its own text and pictures, puts it right after the original, and hides the original (nothing is deleted)."
)
if st.button("Find slides with a broken layout"):
    state.broken = broken_layout_candidates()
ticked = []
for issue in state.get("broken") or []:
    image, text = st.columns([1, 2])
    image.image(slide_image(issue.slide_id), width="stretch")
    text.markdown(f"**Slide {issue.slide_id + 1}** · score {issue.score:.2f}\n\n" + "\n".join(f"- {r}" for r in issue.reasons))
    if text.checkbox("Repair this slide", key=f"repair-{issue.slide_id}"):
        ticked.append(issue.slide_id)
if state.get("broken"):
    from_file = [sid for sid in ticked if use_reviewed and sid in reviewed_repairs]
    live = [sid for sid in ticked if sid not in from_file]
    if st.button(f"Plan repairs for the {len(ticked)} ticked slide(s)", type="primary", disabled=not ticked or (bool(live) and not api_key)):
        checked = [copy.deepcopy(c) for sid in from_file for c in reviewed_repairs[sid]]
        if live:
            try:
                _, _, model, repairer = llm(api_key)
            except Exception as e:
                st.error(f"Could not connect to Groq: {e}")
                st.stop()
            with st.status(f"Rebuilding {len(live)} slide(s) with {model}…", expanded=True) as status:
                repairs = repair.plan_repairs(repairer, SOURCE_PPTX, live, deck)
                checked += editor.check_edits(repairs, dossiers, deck)
                status.update(label=f"{sum(c.accepted for c in checked)} of {len(ticked)} repair(s) passed the checks", state="complete")
        state.suggestions[REPAIRS] = checked
        for c in checked:
            state.approved[edit_key(c.edit)] = c.accepted
    if live and not api_key and ticked:
        st.caption("Some ticked slides have no pre-checked repair; a Groq key is needed to repair them live.")

def show_repair(c) -> None:
    e = c.edit
    st.subheader(f"Repair of slide {e.after_slide_id + 1}")
    before, after = st.columns(2)
    before.caption("Now (broken)")
    before.image(slide_image(e.after_slide_id), width="stretch")
    after.caption("Rebuilt (drawn from the real rebuilt slide; open the deck for exact fonts)")
    try:
        after.image(rebuilt_preview(e), width="stretch")
    except Exception as err:  # the preview is a convenience; never block the review on it
        after.caption(f"(No preview available: {err})")
    keep = st.checkbox(
        "Keep the slide's pictures",
        value=e.keep_pictures,
        key=f"pictures-{edit_key(e)}",
        help="Untick when what is left of the pictures is noise (stacked animation frames, a lone logo); the slide is then rebuilt as text only.",
    )
    if keep != e.keep_pictures:
        e.keep_pictures = keep
        st.rerun()
    with st.expander("Rebuilt text"):
        st.markdown(f"**{e.title}**\n\n" + "\n".join(f"- {b}" for b in e.bullets))


repairs_shown = state.suggestions.get(REPAIRS, [])
for c in [c for c in repairs_shown if c.accepted]:
    show_repair(c)
    st.success("Uses only the slide's own words and keeps most of them. Check the order and meaning before approving.")
    key = edit_key(c.edit)
    state.approved[key] = st.checkbox("Approve this repair", value=state.approved.get(key, True), key=f"approve-{key}")
    st.divider()
held_back = [c for c in repairs_shown if not c.accepted]
if held_back:
    with st.expander(f"Held back by the review ({len(held_back)}) — never applied unless you approve them"):
        for c in held_back:
            show_repair(c)
            st.caption("Why it was held back: " + " · ".join(c.problems))
            key = edit_key(c.edit)
            state.approved[key] = st.checkbox("Approve anyway", value=state.approved.get(key, False), key=f"approve-{key}")
            st.divider()
if state.suggestions.get(REPAIRS):
    st.info("Approved repairs are applied with the other approved changes: use “Apply approved slides” above.")
